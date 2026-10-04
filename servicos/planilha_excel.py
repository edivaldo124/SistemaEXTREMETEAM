"""Planilha econômica em Excel (.xlsx), montada com openpyxl.

Quatro abas: o resumo mês a mês, as entradas e as saídas linha a linha, e as saídas por
categoria em cada mês. Tudo vai como valor, sem fórmula: o openpyxl não calcula fórmula,
e quem abre o arquivo no celular ou na prévia do e-mail veria as células de total vazias.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from servicos import planilha
from servicos.formatacao import formatar_competencia, rotulo_forma_pagamento

FORMATO_MOEDA = '"R$" #,##0.00;[Red]-"R$" #,##0.00'
FORMATO_DATA = 'DD/MM/YYYY'
MESES_CURTOS = ('Jan', 'Fev', 'Mar', 'Abr', 'Mai', 'Jun', 'Jul', 'Ago', 'Set', 'Out', 'Nov', 'Dez')

# A fonte vai nomeada: sem nome, o LibreOffice cai numa serifada só nas células com estilo.
_CABECALHO_FONTE = Font(name='Calibri', bold=True, color='F5B914')
_CABECALHO_FUNDO = PatternFill('solid', fgColor='000000')
_TOTAL_FONTE = Font(name='Calibri', bold=True)
_TOTAL_BORDA = Border(top=Side(style='thin', color='000000'))


def _escrever(aba, linha, valores, formatos=None, *, total=False):
    for coluna, valor in enumerate(valores, start=1):
        celula = aba.cell(row=linha, column=coluna, value=valor)
        if isinstance(valor, str):
            # Texto que começa com "=" viraria fórmula no openpyxl: uma despesa chamada
            # "=HYPERLINK(...)" executaria ao abrir. Fica sempre como texto.
            celula.data_type = 's'
        formato = (formatos or {}).get(coluna)
        if formato:
            celula.number_format = formato
        if total:
            celula.font = _TOTAL_FONTE
            celula.border = _TOTAL_BORDA


def _cabecalho(aba, titulos, larguras, direita=()):
    """Cabeçalho preto com letra dourada; `direita` são as colunas de valor, alinhadas
    como os números embaixo delas."""
    _escrever(aba, 1, titulos)
    for coluna, largura in enumerate(larguras, start=1):
        celula = aba.cell(row=1, column=coluna)
        celula.font = _CABECALHO_FONTE
        celula.fill = _CABECALHO_FUNDO
        celula.alignment = Alignment(vertical='center', horizontal='right' if coluna in direita else None)
        aba.column_dimensions[get_column_letter(coluna)].width = largura
    aba.freeze_panes = 'A2'
    # Na impressão, todas as colunas numa folha deitada; as linhas seguem nas próximas.
    aba.page_setup.orientation = 'landscape'
    aba.sheet_properties.pageSetUpPr.fitToPage = True
    aba.page_setup.fitToWidth = 1
    aba.page_setup.fitToHeight = 0


def _aba_resumo(aba, resumo):
    _cabecalho(aba, ['Mês', 'Entradas', 'Saídas', 'Saldo do mês', 'Acumulado no ano'], [16, 16, 16, 16, 18], direita=range(2, 6))
    moeda = {c: FORMATO_MOEDA for c in range(2, 6)}
    linha = 2
    for mes in resumo.meses:
        _escrever(aba, linha, [mes.nome, mes.entradas, mes.saidas, mes.saldo, mes.acumulado], moeda)
        linha += 1
    _escrever(aba, linha, ['Total do ano', resumo.entradas, resumo.saidas, resumo.saldo, None], moeda, total=True)


def _aba_entradas(aba, pagamentos):
    _cabecalho(aba, ['Data do pagamento', 'Aluno', 'Plano', 'Competência', 'Forma de pagamento', 'Valor'],
               [18, 34, 22, 20, 22, 14], direita=(6,))
    formatos = {1: FORMATO_DATA, 6: FORMATO_MOEDA}
    for linha, pagamento in enumerate(pagamentos, start=2):
        _escrever(aba, linha, [
            pagamento.data_pagamento or pagamento.vencimento,
            pagamento.aluno.nome if pagamento.aluno else '—',
            pagamento.plano.nome_plano if pagamento.plano else '—',
            formatar_competencia(pagamento.competencia),
            rotulo_forma_pagamento(pagamento.forma_pagamento),
            pagamento.valor,
        ], formatos)
    if pagamentos:
        aba.auto_filter.ref = f'A1:F{len(pagamentos) + 1}'


def _aba_saidas(aba, despesas):
    _cabecalho(aba, ['Data', 'Categoria', 'Descrição', 'Valor', 'Lançada por'], [14, 24, 44, 14, 18], direita=(4,))
    formatos = {1: FORMATO_DATA, 4: FORMATO_MOEDA}
    for linha, despesa in enumerate(despesas, start=2):
        _escrever(aba, linha, [
            despesa.data, despesa.rotulo_categoria, despesa.descricao, despesa.valor, despesa.lancado_por or '—',
        ], formatos)
    if despesas:
        aba.auto_filter.ref = f'A1:E{len(despesas) + 1}'


def _aba_categorias(aba, resumo):
    _cabecalho(aba, ['Categoria', *MESES_CURTOS, 'Total'], [24] + [12] * 12 + [14], direita=range(2, 15))
    moeda = {c: FORMATO_MOEDA for c in range(2, 15)}
    linha = 2
    for categoria in resumo.categorias:
        _escrever(aba, linha, [
            categoria.rotulo,
            *[categoria.por_mes.get(numero) for numero in range(1, 13)],
            categoria.total,
        ], moeda)
        linha += 1
    _escrever(aba, linha, [
        'Total',
        *[mes.saidas if mes else None for mes in map(resumo.mes, range(1, 13))],
        resumo.saidas,
    ], moeda, total=True)


def gerar(ano, *, hoje):
    """Bytes do arquivo .xlsx do ano."""
    resumo = planilha.resumo_anual(ano, hoje=hoje)

    livro = Workbook()
    livro.properties.creator = 'Extreme Team'
    livro.properties.title = f'Planilha econômica {ano}'

    _aba_resumo(livro.active, resumo)
    livro.active.title = f'Resumo {ano}'
    _aba_entradas(livro.create_sheet('Entradas'), planilha.entradas_detalhadas(ano))
    _aba_saidas(livro.create_sheet('Saídas'), planilha.saidas_detalhadas(ano))
    _aba_categorias(livro.create_sheet('Por categoria'), resumo)

    saida = BytesIO()
    livro.save(saida)
    return saida.getvalue()
