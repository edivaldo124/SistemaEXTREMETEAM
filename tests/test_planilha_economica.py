"""Planilha econômica: fluxo de caixa, lançamento de despesas e download em Excel."""
import os
from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from dao.despesaDAO import DespesaDAO
from modelos.despesa import Despesa
from servicos import planilha


@pytest.fixture
def criar_despesa(contexto_app):
    def _criar(**overrides):
        dados = dict(data=date.today(), categoria='aluguel', descricao='Aluguel do salão', valor=Decimal('1000.00'))
        dados.update(overrides)
        despesa = Despesa(**dados)
        DespesaDAO.salvar(despesa)
        return despesa
    return _criar


def _form(**overrides):
    dados = {'data': date.today().isoformat(), 'categoria': 'contas', 'descricao': 'Conta de luz', 'valor': '320,50'}
    dados.update(overrides)
    return dados


# ------------------------------------------------------------------ cálculo --

def test_entrada_conta_no_mes_em_que_foi_paga(criar_pagamento, contexto_app):
    # Vencia em janeiro, foi paga em fevereiro: o dinheiro entrou em fevereiro.
    criar_pagamento(valor=150, vencimento=date(2025, 1, 10), status='pago', data_pagamento=date(2025, 2, 3))
    # Lançamento antigo sem data de pagamento: cai no mês do vencimento.
    criar_pagamento(valor=100, vencimento=date(2025, 3, 5), status='pago', data_pagamento=None)

    resumo = planilha.resumo_anual(2025, hoje=date(2025, 12, 31))

    assert resumo.mes(1).entradas == Decimal('0.00')
    assert resumo.mes(2).entradas == Decimal('150.00')
    assert resumo.mes(3).entradas == Decimal('100.00')
    assert resumo.entradas == Decimal('250.00')


def test_so_mensalidade_paga_e_entrada(criar_pagamento, contexto_app):
    for status in ('pendente', 'atrasado', 'em_analise', 'recusado', 'cancelado', 'reembolsado'):
        criar_pagamento(valor=90, vencimento=date(2025, 5, 10), status=status, data_pagamento=date(2025, 5, 10))
    criar_pagamento(valor=150, vencimento=date(2025, 5, 10), status='pago', data_pagamento=date(2025, 5, 10))

    resumo = planilha.resumo_anual(2025, hoje=date(2025, 12, 31))

    assert resumo.entradas == Decimal('150.00')


def test_saldo_e_acumulado_mes_a_mes(criar_pagamento, criar_despesa):
    criar_pagamento(valor=500, vencimento=date(2025, 1, 5), status='pago', data_pagamento=date(2025, 1, 5))
    criar_despesa(data=date(2025, 1, 20), valor=Decimal('800.00'))
    criar_pagamento(valor=1000, vencimento=date(2025, 2, 5), status='pago', data_pagamento=date(2025, 2, 5))
    criar_despesa(data=date(2025, 2, 20), valor=Decimal('100.00'))

    resumo = planilha.resumo_anual(2025, hoje=date(2025, 12, 31))

    janeiro, fevereiro = resumo.mes(1), resumo.mes(2)
    assert (janeiro.saldo, janeiro.acumulado) == (Decimal('-300.00'), Decimal('-300.00'))
    assert (fevereiro.saldo, fevereiro.acumulado) == (Decimal('900.00'), Decimal('600.00'))
    assert resumo.mes(12).acumulado == Decimal('600.00')
    assert (resumo.entradas, resumo.saidas, resumo.saldo) == (Decimal('1500.00'), Decimal('900.00'), Decimal('600.00'))


def test_mes_que_nao_comecou_so_aparece_com_movimento(criar_despesa):
    criar_despesa(data=date(2026, 12, 10), valor=Decimal('50.00'))

    resumo = planilha.resumo_anual(2026, hoje=date(2026, 10, 4))

    assert [m.numero for m in resumo.meses] == list(range(1, 11)) + [12]
    assert planilha.resumo_anual(2027, hoje=date(2026, 10, 4)).meses == []


def test_saidas_por_categoria_maior_primeiro(criar_despesa):
    criar_despesa(data=date(2025, 3, 1), categoria='contas', valor=Decimal('250.00'))
    criar_despesa(data=date(2025, 4, 1), categoria='contas', valor=Decimal('250.00'))
    criar_despesa(data=date(2025, 3, 1), categoria='aluguel', valor=Decimal('1500.00'))
    criar_despesa(data=date(2024, 3, 1), categoria='marketing', valor=Decimal('999.00'))  # outro ano

    resumo = planilha.resumo_anual(2025, hoje=date(2025, 12, 31))

    assert [(c.chave, c.total, c.percentual) for c in resumo.categorias] == [
        ('aluguel', Decimal('1500.00'), 75),
        ('contas', Decimal('500.00'), 25),
    ]
    assert resumo.categorias[1].por_mes == {3: Decimal('250.00'), 4: Decimal('250.00')}


def test_anos_disponiveis_vao_do_primeiro_movimento_ate_hoje(criar_pagamento, criar_despesa):
    criar_pagamento(valor=150, vencimento=date(2023, 6, 1), status='pago', data_pagamento=date(2023, 6, 1))
    criar_pagamento(valor=150, vencimento=date(2019, 6, 1), status='pendente')  # não pago: não conta
    criar_despesa(data=date(2024, 2, 1))

    assert planilha.anos_disponiveis(hoje=date(2026, 10, 4)) == [2026, 2025, 2024, 2023]


# ------------------------------------------------------------------- acesso --

@pytest.mark.parametrize('metodo, url', [
    ('get', '/admin/planilha'),
    ('get', f'/admin/planilha/{date.today().year}.xlsx'),
    ('post', '/admin/planilha/despesas'),
    ('post', '/admin/planilha/despesas/1'),
    ('post', '/admin/planilha/despesas/1/excluir'),
])
def test_planilha_e_so_do_admin(client, criar_aluno, logar_como_aluno, metodo, url):
    resposta = getattr(client, metodo)(url, data=_form() if metodo == 'post' else None)
    assert resposta.status_code == 302 and resposta.headers['Location'].endswith('/login')

    logar_como_aluno(criar_aluno())
    resposta = getattr(client, metodo)(url, data=_form() if metodo == 'post' else None)
    assert resposta.status_code == 302 and resposta.headers['Location'].endswith('/login')


# --------------------------------------------------------------------- tela --

def test_tela_mostra_totais_e_saidas_do_mes(client, logar_como_admin, criar_pagamento, criar_despesa):
    hoje = date.today()
    criar_pagamento(valor=2000, vencimento=hoje, status='pago', data_pagamento=hoje)
    criar_despesa(data=hoje, descricao='Aluguel de outubro', valor=Decimal('1234.56'))
    logar_como_admin()

    resposta = client.get('/admin/planilha')

    html = resposta.get_data(as_text=True)
    assert resposta.status_code == 200
    assert 'Planilha econômica' in html
    assert 'R$ 2.000,00' in html and 'R$ 1.234,56' in html and 'R$ 765,44' in html
    assert 'Aluguel de outubro' in html
    assert f'/admin/planilha/{hoje.year}.xlsx' in html
    assert 'aria-current="page"' in html and 'href="/admin/planilha"' in html


def test_ano_sem_movimento_mostra_estado_vazio(client, logar_como_admin):
    logar_como_admin()
    proximo = date.today().year + 1

    html = client.get(f'/admin/planilha?ano={proximo}').get_data(as_text=True)

    assert f'Nenhum movimento em {proximo}' in html


def test_ano_fora_do_intervalo_volta_ao_atual(client, logar_como_admin):
    logar_como_admin()

    html = client.get('/admin/planilha?ano=1850&mes=99').get_data(as_text=True)

    assert f'<option value="{date.today().year}" selected>' in html


# ----------------------------------------------------------------- despesas --

def test_lancar_despesa(client, logar_como_admin, contexto_app):
    logar_como_admin()
    data = date.today().replace(day=1)

    resposta = client.post('/admin/planilha/despesas', data=_form(data=data.isoformat(), valor='1.234,56'))

    assert resposta.status_code == 302
    assert f'ano={data.year}' in resposta.headers['Location'] and f'mes={data.month}' in resposta.headers['Location']
    despesa = Despesa.query.one()
    assert (despesa.data, despesa.categoria, despesa.descricao, despesa.valor) == (
        data, 'contas', 'Conta de luz', Decimal('1234.56'),
    )
    assert despesa.lancado_por == os.environ['ADMIN_USER']


@pytest.mark.parametrize('campo, valor, mensagem', [
    ('data', '', 'Informe a data da despesa.'),
    ('data', '1999-12-31', 'Use uma data entre 2000'),
    ('categoria', 'cerveja', 'Escolha uma categoria.'),
    ('descricao', '   ', 'Descreva a despesa.'),
    ('descricao', 'x' * 121, 'Use até 120 caracteres.'),
    ('valor', '0', 'Informe um valor maior que zero'),
    ('valor', '-10', 'Informe um valor maior que zero'),
    ('valor', 'abc', 'Informe um valor maior que zero'),
    ('valor', 'NaN', 'Informe um valor maior que zero'),
    ('valor', '100000000', 'Informe um valor maior que zero'),
])
def test_despesa_invalida_nao_grava_e_aponta_o_campo(client, logar_como_admin, contexto_app, campo, valor, mensagem):
    logar_como_admin()

    resposta = client.post('/admin/planilha/despesas', data=_form(**{campo: valor}))

    html = resposta.get_data(as_text=True)
    assert resposta.status_code == 400
    assert mensagem in html
    assert f'aria-describedby="despesa-{campo}-erro"' in html
    assert Despesa.query.count() == 0
    if campo != 'descricao':
        assert 'value="Conta de luz"' in html  # o resto do que foi digitado volta preenchido


def test_alterar_despesa(client, logar_como_admin, criar_despesa):
    despesa = criar_despesa(descricao='Luz', valor=Decimal('10.00'))
    logar_como_admin()

    form = client.get(f'/admin/planilha?editar={despesa.id}').get_data(as_text=True)
    assert 'Salvar alteração' in form and 'value="10,00"' in form

    resposta = client.post(f'/admin/planilha/despesas/{despesa.id}',
                           data=_form(descricao='Luz de setembro', valor='99,90', categoria='contas'))

    assert resposta.status_code == 302
    alterada = DespesaDAO.buscar_por_id(despesa.id)
    assert (alterada.descricao, alterada.valor, alterada.categoria) == ('Luz de setembro', Decimal('99.90'), 'contas')


def test_alterar_com_erro_continua_editando(client, logar_como_admin, criar_despesa):
    despesa = criar_despesa(descricao='Luz')
    logar_como_admin()

    resposta = client.post(f'/admin/planilha/despesas/{despesa.id}', data=_form(valor=''))

    assert resposta.status_code == 400
    assert 'Salvar alteração' in resposta.get_data(as_text=True)
    assert DespesaDAO.buscar_por_id(despesa.id).descricao == 'Luz'


def test_excluir_despesa(client, logar_como_admin, criar_despesa):
    despesa = criar_despesa()
    logar_como_admin()

    resposta = client.post(f'/admin/planilha/despesas/{despesa.id}/excluir')

    assert resposta.status_code == 302
    assert Despesa.query.count() == 0


@pytest.mark.parametrize('url', ['/admin/planilha/despesas/999', '/admin/planilha/despesas/999/excluir'])
def test_despesa_inexistente_e_404(client, logar_como_admin, contexto_app, url):
    logar_como_admin()
    assert client.post(url, data=_form()).status_code == 404


# -------------------------------------------------------------------- excel --

def test_excel_do_ano(client, logar_como_admin, criar_aluno, criar_pagamento, criar_despesa):
    ano = date.today().year
    aluno = criar_aluno(nome='Maria Lutadora')
    criar_pagamento(aluno=aluno, valor=150, vencimento=date(ano, 1, 10), status='pago',
                    data_pagamento=date(ano, 1, 12), competencia=f'{ano}-01', forma_pagamento='pix')
    criar_despesa(data=date(ano, 1, 20), categoria='contas', descricao='=HYPERLINK("http://x","clique")',
                  valor=Decimal('40.00'))
    logar_como_admin()

    resposta = client.get(f'/admin/planilha/{ano}.xlsx')

    assert resposta.status_code == 200
    assert resposta.mimetype == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    assert f'filename="planilha-economica-{ano}.xlsx"' in resposta.headers['Content-Disposition']
    assert 'no-store' in resposta.headers['Cache-Control']

    livro = load_workbook(BytesIO(resposta.data))
    assert livro.sheetnames == [f'Resumo {ano}', 'Entradas', 'Saídas', 'Por categoria']

    resumo = livro[f'Resumo {ano}']
    assert [c.value for c in resumo[2]] == ['Janeiro', 150, 40, 110, 110]
    total = next(linha for linha in resumo.iter_rows(values_only=True) if linha[0] == 'Total do ano')
    assert total[1:4] == (150, 40, 110)

    entradas = livro['Entradas']
    assert [c.value for c in entradas[2]][1:] == ['Maria Lutadora', 'Mensal', f'Janeiro de {ano}', 'Pix', 150]
    assert entradas['A2'].value.date() == date(ano, 1, 12)

    # Texto que parece fórmula fica como texto: nada executa ao abrir o arquivo.
    descricao = livro['Saídas']['C2']
    assert descricao.value == '=HYPERLINK("http://x","clique")'
    assert descricao.data_type == 's'

    categorias = livro['Por categoria']
    assert [c.value for c in categorias[2]][:2] == ['Água, luz e internet', 40]


def test_excel_de_ano_invalido_e_404(client, logar_como_admin):
    logar_como_admin()
    assert client.get('/admin/planilha/1850.xlsx').status_code == 404
